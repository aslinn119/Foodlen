"""Create 李馨瑶's standalone Task 1 report chapter from measured run artifacts."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
OUT = BASE / "task1_李馨瑶.docx"
FIG = BASE / "report_figures" / "task1"
RUN = BASE / "v2_attention" / "runs"
RESULTS = BASE / "results"


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def set_cell_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag); borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc; tc_pr = tc.get_or_add_tcPr(); margins = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}"); node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa"); margins.append(node)
    tc_pr.append(margins)


def set_run_font(run, name="宋体", size=10.5, bold=False, color="000000"):
    run.font.name = name; run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", style=None, align=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        a=p.add_run(bold_prefix);set_run_font(a,bold=True)
        b=p.add_run(text[len(bold_prefix):]);set_run_font(b)
    else:
        r=p.add_run(text);set_run_font(r)
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.25
    return p


def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f"Heading {level}")
    r=p.add_run(text);set_run_font(r,"黑体",15 if level==1 else 12,bold=True)
    p.paragraph_format.space_before=Pt(11);p.paragraph_format.space_after=Pt(5)
    return p


def add_caption(doc, text):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text);set_run_font(r,"宋体",9.5)
    p.paragraph_format.space_after=Pt(7)


def add_image(doc, path, width=6.0):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Inches(width))


def make_table(doc, headers, rows, widths=None):
    table=doc.add_table(rows=1,cols=len(headers));table.autofit=False
    table.style="Table Grid"
    for i,h in enumerate(headers):
        c=table.rows[0].cells[i];c.text=str(h);set_cell_shading(c,"FFFFFF");set_cell_borders(c);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;set_cell_margins(c)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:set_run_font(r,"宋体",9,bold=True,color="000000")
    for ridx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val);set_cell_shading(cells[i],"FFFFFF");set_cell_borders(cells[i]);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i>0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:set_run_font(r,"宋体",9)
    if widths:
        for row in table.rows:
            for i,w in enumerate(widths):row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return table


def scores(path):
    d=read_json(path)
    if not d:return None
    vals=d["bleu1_to_4"]
    return [f"{x:.2f}" for x in vals],d.get("n","")


def extra_metrics(path):
    d=read_json(path)
    return d if d and "rouge_l" in d and "cider_d" in d else None


def main():
    history=read_json(RUN/"v2"/"history.json") or []
    config=read_json(RUN/"v2"/"config.json") or {}
    ablation_history=read_json(RUN/"v1_ablation"/"history.json") or []
    v1=scores(RESULTS/"v1_ablation_test.json")
    v2=scores(RESULTS/"v2_greedy_test.json")
    beam=scores(RESULTS/"v2_beam_test.json")
    m_v1=extra_metrics(RESULTS/"v1_ablation_test.json")
    m_v2=extra_metrics(RESULTS/"v2_greedy_test.json")
    m_beam=extra_metrics(RESULTS/"v2_beam_test.json")
    v1_rows=read_json(RESULTS/"v1_ablation_test.json")
    v2_rows=read_json(RESULTS/"v2_greedy_test.json")
    best=max(history,key=lambda x:x["val_bleu4"]) if history else None

    doc=Document()
    sec=doc.sections[0];sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.82);sec.right_margin=Inches(.82)
    normal=doc.styles["Normal"];normal.font.name="宋体";normal._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体");normal.font.size=Pt(10.5)
    for sty in ("Heading 1","Heading 2","Title"):
        doc.styles[sty].font.color.rgb=RGBColor(0,0,0)
        doc.styles[sty]._element.rPr.rFonts.set(qn("w:eastAsia"),"黑体")
    title=doc.add_paragraph(style="Title");title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=title.add_run("FoodLens Task 1 V2 注意力模型实验报告");set_run_font(r,"黑体",18,bold=True)
    subtitle=doc.add_paragraph();subtitle.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("神经网络与深度学习课程设计  李馨瑶"),"宋体",10.5)
    paragraph(doc,"本节围绕食物图像描述中的空间细节建模，介绍在 ResNet18–GRU 基线上加入 Bahdanau 注意力的 V2 模型，并报告真实 Food500Cap 测试集上的对比、消融与解码实验。实验使用同一份清洗后数据划分和训练集词表；注意力权重还用于展示模型在生成不同词时关注的图像区域。",align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    heading(doc,"3.2 改进模型 V2 加入注意力机制",1)
    heading(doc,"3.2.1 设计动机与模型结构",2)
    paragraph(doc,"V1 将 ResNet18 最后一层卷积特征经过全局平均池化压缩为一个 512 维向量，再由 GRU 逐词生成描述。单个全局向量难以保留菜品不同位置的局部信息。对于包含主食、配菜和装饰物的复杂图像，解码器可能无法在生成食材词时回到对应区域。V2 保留 V1 的 ResNet18 和 GRU 主干，将池化前特征图保留为 7×7 网格，使编码器输出 49 个 512 维空间特征。")
    paragraph(doc,"在时间步 t，注意力模块把上一步隐状态 h(t−1) 与第 i 个空间特征 a(i) 投影到共同空间，并计算兼容度 e(t,i)=vᵀtanh(W h(t−1)+U a(i))。对 49 个位置做 softmax 得到权重 α(t,i)，再按权重加权求和得到上下文 c(t)=Σᵢ α(t,i)a(i)。GRU 的输入由当前词嵌入与 c(t) 拼接构成，输出层预测下一个词。由于注意力依赖每一步的隐状态，同一张图在生成不同词时可以采用不同空间权重。")
    paragraph(doc,"贪心解码与束搜索共享同一编码器和注意力解码器。贪心策略每步选择概率最大的词；束搜索保留累计对数概率较高的多个候选序列，并按长度惩罚调整分数。解码时记录每个生成词对应的 49 维注意力权重，插值到 224×224 后叠加在模型实际输入裁剪图上。EOS 终止词不参与词与热图对齐。")
    add_image(doc,FIG/"v2_architecture.png",6.1);add_caption(doc,"图 3  V2 空间注意力图像描述模型结构")

    heading(doc,"3.2.2 数据与训练设置",2)
    paragraph(doc,"清洗数据包含 train 17,784 张、val 1,967 张和 test 4,938 张图像，共 24,689 张；数据审核未发现缺失图像、重复路径或空 caption。数据说明记录了对 126 张 train/test 重复图的去重、按训练集内部划分验证集及 494 个规范类别。图像使用与 V1 相同的 Resize 256、224 裁剪和 ImageNet 均值方差归一化；训练阶段随机裁剪和水平翻转，验证与测试采用中心裁剪。词表从训练 caption 按空格切词并以 min_count=5 重建，包含 2,681 个 token。压缩包未包含说明文档中提到的 vocab.json，因此本实验将重建规则和生成文件纳入交付。")
    make_table(doc,["设置项","取值"],[
        ["图像编码器","ImageNet 预训练 ResNet18，输出 7×7×512 空间特征"],
        ["解码器","Bahdanau 加性注意力 + 单层 GRU"],
        ["Embedding / hidden","256 / 512"],
        ["Batch / 最大长度","64 / 30"],
        ["优化与学习率","Adam；encoder 1×10⁻⁴，decoder 1×10⁻³"],
        ["正则与优化稳定","dropout 0.5，梯度裁剪 5，ReduceLROnPlateau"],
        ["模型选择","验证集前 500 条的 BLEU-4，patience=4，seed=42"],
    ],[1.7,4.6])
    paragraph(doc,"V2 的训练损失采用忽略 PAD 的交叉熵，优化器、预训练编码器和主要超参数沿用组内 V1 源码默认值。训练集词表与 split 的哈希、软硬件版本、参数和逐轮日志均保存在运行目录中。先在真实数据固定 batch 上完成 20 步 smoke test，损失由 7.981 降至 3.812，确认数据管线和梯度更新可用后再进行完整训练。")

    heading(doc,"5.1 Task 1 结果与分析补充",1)
    heading(doc,"5.1.1 训练过程与测试集指标",2)
    if history and ablation_history and (FIG/"training_curves.png").exists():
        add_image(doc,FIG/"training_curves.png",6.2);add_caption(doc,"图 7  V1 无注意力消融与 V2 的训练损失及验证 BLEU-4")
    else:
        paragraph(doc,"训练曲线将在 V1 消融实验完成后更新。")
    rows=[]
    legacy_info=read_json(BASE/"baseline_reference"/"group_v1_reported_metrics.json")
    legacy=legacy_info["bleu1_to_4"] if legacy_info else [30.8849,15.2626,8.2645,4.5702]
    rows.append(["组长提供的 V1 历史结果","贪心","4,938",*[f"{x:.2f}" for x in legacy]])
    for label,result,decode in (("V1 无注意力消融",v1,"贪心"),("V2 注意力",v2,"贪心"),("V2 注意力",beam,"Beam-3")):
        rows.append([label,decode,result[1] if result else "待评测",*(result[0] if result else ["—"]*4)])
    make_table(doc,["模型","解码","样本数","BLEU-1","BLEU-2","BLEU-3","BLEU-4"],rows,[1.42,.72,.72,.73,.73,.73,.73])
    if (FIG/"test_bleu_comparison.png").exists():
        add_image(doc,FIG/"test_bleu_comparison.png",6.15);add_caption(doc,"图 8  完整测试集 BLEU-1 至 BLEU-4 对比")
    metric_rows=[["组长提供的 V1 历史结果","—","—"],
                 ["V1 无注意力消融",f"{m_v1['rouge_l']:.4f}" if m_v1 else "待评测",f"{m_v1['cider_d']:.3f}" if m_v1 else "待评测"],
                 ["V2 注意力 greedy",f"{m_v2['rouge_l']:.4f}" if m_v2 else "待评测",f"{m_v2['cider_d']:.3f}" if m_v2 else "待评测"],
                 ["V2 注意力 beam-3",f"{m_beam['rouge_l']:.4f}" if m_beam else "待评测",f"{m_beam['cider_d']:.3f}" if m_beam else "待评测"]]
    make_table(doc,["模型","ROUGE-L","CIDEr-D"],metric_rows,[2.8,1.7,1.7])
    paragraph(doc,"本实验报告 BLEU-1 至 BLEU-4、ROUGE-L 和 CIDEr-D。ROUGE-L 按 token 级最长公共子序列及 β=1.2 汇总；CIDEr-D 使用 1–4 gram TF-IDF、σ=6 的长度惩罚和 ×10 标度。当前 Food500Cap 标注每张图只有一条参考描述，因此这些分数均按单参考计算，和多参考评测结果不可直接横向比较。")
    if best:
        paragraph(doc,f"V2 共完成 {len(history)} 轮训练；第 {best['epoch']} 轮在验证集前 500 条上取得最高 BLEU-4={best['val_bleu4']:.2f}，该轮模型用于最终测试。测试集 BLEU 由组内复用的语料级 bleu.py 对完整 test split 计算。历史 V1 行来自组长提供的记录，源码压缩包未附与该数值对应的模型文件，因此该行无法独立复现；V1 无注意力消融则是在本目录使用同一数据与训练设置重新训练的结果。")
    else:
        paragraph(doc,"模型选择与测试集数值将在正式训练、评测后由日志自动写入。")
    heading(doc,"5.1.2 注意力消融与解码方式",2)
    if v2 and legacy:
        delta0=float(v2[0][3])-float(legacy[3])
        rel=delta0/float(legacy[3])*100
        paragraph(doc,f"与组长提供的 V1 历史记录相比，V2 greedy 的 BLEU-4 高 {delta0:.2f} 分（相对提升 {rel:.1f}%）。由于旧 checkpoint 缺失，这一差异只能作为同划分、同 BLEU 口径下的报告值对照，不能排除训练配置或随机种子的影响。")
    if v1 and v2 and beam:
        delta=float(v2[0][3])-float(v1[0][3])
        bdelta=float(beam[0][3])-float(v2[0][3])
        paragraph(doc,f"在重新训练的 V1 结构与 V2 之间，BLEU-4 变化为 {delta:+.2f}；该消融用于估计保留空间特征和注意力的整体影响。V2 Beam-3 相对贪心的 BLEU-4 变化为 {bdelta:+.2f}。这两个比较均在同一 test split 和同一评价实现下计算；若 beam 未提升，说明累计序列概率及长度惩罚并不必然与单参考 BLEU 一致。")
    else:
        paragraph(doc,"无注意力消融和 V2 Beam-3 完整测试集结果生成后，本段将据实补充。")
    paragraph(doc,"注意力图提供的是模型内部空间权重而非因果解释。可视化用于观察生成词与高权重区域是否大致对应，不能单独证明模型正确识别了食材。")

    heading(doc,"5.1.3 注意力可视化与误差分析",2)
    image_files=sorted((FIG/"attention_samples").glob("*/attention.png")) if (FIG/"attention_samples").exists() else []
    if v1_rows and v2_rows and image_files:
        p1={r["file_name"]:r for r in v1_rows["predictions"]};p2={r["file_name"]:r for r in v2_rows["predictions"]}
        for path in image_files[:3]:
            meta=read_json(path.parent/"caption.json") or {};name=meta.get("image","")
            normalized=name.replace("\\","/")
            rel=normalized.split("Food500Cap_ready/",1)[-1] if "Food500Cap_ready/" in normalized else normalized
            if rel in p1 and rel in p2:
                paragraph(doc,f"{meta.get('category','样例')}（{rel}）参考描述：{p2[rel]['reference']}")
                paragraph(doc,f"V1 无注意力：{p1[rel]['hypothesis']}；V2 注意力：{p2[rel]['hypothesis']}")
    for i,path in enumerate(image_files[:3],1):
        meta=read_json(path.parent/"caption.json") or {}
        paragraph(doc,f"样例 {i}：模型生成“{meta.get('caption','')}”。热图按生成词顺序展示注意力分布；比较菜品主体与配料位置可观察空间权重的变化。参考描述：{meta.get('reference','待补充')}")
        add_image(doc,path,5.6);add_caption(doc,f"图 9-{i}  测试样例 {i} 的词级注意力热图")
    if not image_files:
        paragraph(doc,"测试集代表样例的词级注意力图将在最终 checkpoint 生成后插入。")
    else:
        paragraph(doc,"样例中，Arab salad 的生成词覆盖了番茄、黄瓜和 parsley，热图在不同词之间发生变化；Chinese chicken salad 虽识别菜名和鸡肉，却重复生成 chicken 并漏掉杏仁、芝麻等配料。番茄汤样例生成了 soup、basil 和 spoon 等相关词，但部分高权重落在碗沿和画面边缘。可见注意力为局部描述提供了可视化线索，但空间权重并非总与目标食材精确重合，模型仍有重复、遗漏和背景关注问题。")
    paragraph(doc,"Food-500 Cap 的 494 个菜品类别具有细粒度外观差异，描述又以单条参考文本进行 BLEU 评价，语义正确但用词不同也可能得分偏低。当前模型仍可能生成模板化句式、漏掉低频食材或过早结束。后续可在保持数据划分不变的前提下比较更大解码器、scheduled sampling、多参考评价和更强视觉编码器，并以验证集选择模型，避免用 test 集调参。")

    doc.save(OUT)
    print(OUT)


if __name__=="__main__":main()
